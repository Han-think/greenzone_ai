"""5단계: 결과물 만들기 (outputs/greenzone2nd/)

  greenzone2nd.geojson                   거점 9곳 (웹 지도 메인 레이어)
  greenzone2nd_layer_zones.geojson       생활권 9곳 영역
  greenzone2nd_layer_candidates.geojson  1차 후보 11곳 + 연결된 생활권·역할(거점/예비)
  greenzone2nd_layer_alternatives.geojson 대안 지점(실사 1~3순위) + 생활권 최적 도로 지점(참고)
  greenzone2nd_layer_households.geojson  세대 필지 점 (세대 수, 생활권)
  greenzone2nd.csv / .xlsx               표
  greenzone2nd_map.png / _map.html       지도
  greenzone2nd_sensitivity.png           Louvain 설정별 생활권 수
"""
import json
import logging
import numpy as np
import pandas as pd
import geopandas as gpd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
from shapely.geometry import Point
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from gzlib import *

S = pd.read_pickle(b("plan.pkl"))
zones, cand, ALT, REF, HUB, CMP, hh = S["zones"], S["cand"], S["ALT"], S["REF"], S["HUB"], S["CMP"], S["hh"]
SENS = pd.read_csv(b("louvain_sens.csv"))
META = load_json("louvain_meta.json")
HS = load_json("households_stats.json")
pf = gpd.read_file(PROTECT).to_crs(CRS_M)
D = dong()
NC = len(HUB)


def ll(g):
    p = gpd.GeoSeries([g], crs=CRS_M).to_crs(4326).iloc[0]
    return round(p.y, 6), round(p.x, 6)


# ---- 표 정리 ----------------------------------------------------------------------------------
H = HUB.copy()
H["위도"], H["경도"] = zip(*[ll(g) for g in H.geometry])
H["실사필요"] = H["구분"].str.startswith("대안")
alt1 = ALT[ALT["순위"] == 1].set_index("생활권")
cinfo = cand.set_index("id")
H["접한도로"] = [alt1.at[z, "도로명"] if r.startswith("A") else cinfo.at[r, "접도_도로"] for z, r in zip(H["생활권"], H["거점id"])]
H["도로폭_m"] = [alt1.at[z, "등록도로폭_m"] if r.startswith("A") else cinfo.at[r, "유효폭_m"] for z, r in zip(H["생활권"], H["거점id"])]
H["도로폭_기준"] = ["등록폭" if r.startswith("A") else "유효폭(1차)" for r in H["거점id"]]
H["보호시설까지_m"] = [alt1.at[z, "보호시설까지_m"] if r.startswith("A") else cinfo.at[r, "보호시설과의 거리_m"] for z, r in zip(H["생활권"], H["거점id"])]
H["생활권안_1차후보"] = [", ".join(f"{int(n)}번" for n in cand.loc[cand["생활권"] == z, "번호"]) or "없음" for z in H["생활권"]]
H["생활권_최적도로지점_평균보행_m"] = [REF.set_index("생활권")["생활권_평균보행_m"].get(z, alt1["생활권_평균보행_m"].get(z)) for z in H["생활권"]]
HUBCOLS = ["생활권", "거점id", "구분", "이름", "생활권안_1차후보", "담당세대", "담당필지", "평균보행_m", "세대90%_m", "최대보행_m",
           "200m안", "300m안", "400m안", "접한도로", "도로폭_m", "도로폭_기준", "보호시설까지_m", "생활권_최적도로지점_평균보행_m", "실사필요", "위도", "경도"]
H = H[HUBCOLS + ["geometry"]]

C = cand.copy()
C["위도"], C["경도"] = C["위도"].round(6), C["경도"].round(6)
CCOLS = ["id", "번호", "이름", "부지유형", "생활권", "역할", "생활권_평균보행_m", "생활권_300m안", "생활권_최적도로지점_평균보행_m", "유효폭_m", "보호시설과의 거리_m", "위도", "경도"]
C = C[CCOLS + ["geometry"]]

A = pd.concat([ALT, REF], ignore_index=True)
A["위도"], A["경도"] = zip(*[ll(g) for g in A.geometry])
ACOLS = ["id", "용도", "생활권", "순위", "도로명", "등록도로폭_m", "도로등급", "가까운_건물주소", "최근접_보호시설", "보호시설까지_m", "교차로까지_m",
         "생활권_평균보행_m", "생활권_300m안", "생활권_통과지점수", "위도", "경도"]
A = gpd.GeoDataFrame(A[ACOLS + ["geometry"]], crs=CRS_M)

Z = zones.copy()
Z["세대"] = [int(hh.loc[hh["생활권"] == z, "세대"].sum()) for z in Z["생활권"]]
Z["필지"] = [int((hh["생활권"] == z).sum()) for z in Z["생활권"]]
Z["면적_ha"] = (Z.area / 1e4).round(1)
Z["거점"] = Z["생활권"].map(dict(zip(H["생활권"], H["이름"])))
Z["거점구분"] = Z["생활권"].map(dict(zip(H["생활권"], H["구분"])))
Z = Z[["생활권", "세대", "필지", "면적_ha", "거점", "거점구분", "geometry"]]


def to_geo(g, name, prec=3):
    g = g.copy()
    for c in g.columns:
        if c != "geometry" and g[c].dtype.kind == "f":
            g[c] = g[c].round(prec)
    g.to_crs(4326).to_file(os.path.join(OUT, name), driver="GeoJSON")


to_geo(H, "greenzone2nd.geojson")
to_geo(Z, "greenzone2nd_layer_zones.geojson")
to_geo(C, "greenzone2nd_layer_candidates.geojson")
to_geo(A, "greenzone2nd_layer_alternatives.geojson")
to_geo(hh[["key", "세대", "주택유형", "생활권", "geometry"]], "greenzone2nd_layer_households.geojson")
H.drop(columns="geometry").round(3).to_csv(os.path.join(OUT, "greenzone2nd.csv"), index=False, encoding="utf-8-sig")

# ---- 그림 공통 ----------------------------------------------------------------------------------
logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
for f in font_manager.findSystemFonts():
    try:
        font_manager.fontManager.addfont(f)
    except Exception:
        pass
names = {f.name for f in font_manager.fontManager.ttflist}
plt.rcParams["font.family"] = next((n for n in ("Malgun Gothic", "AppleGothic", "Noto Sans CJK KR", "NanumGothic") if n in names), "sans-serif")
plt.rcParams["axes.unicode_minus"] = False
INK, MUTED, GRID = "#222222", "#666666", "#e5e5e5"
BLUE, ORANGE, GREEN = "#2a78d6", "#eb6834", "#1baf7a"
ZPAL = ["#cfe0f7", "#fbd9c4", "#cdeedd", "#e6d6f2", "#f7d0d6", "#e3ebc4", "#cdebf0", "#f2e2c4", "#d8dcf5"]
ZDOT = ["#6b9fe0", "#ef9a6b", "#5cc49a", "#b38bd6", "#e6849a", "#a8bb5a", "#5fb8c9", "#d6a85c", "#8a95dd"]

# ---- 지도 png ------------------------------------------------------------------------------------
fig, ax = plt.subplots(figsize=(10, 11.5))
for k, r in Z.iterrows():
    gpd.GeoSeries([r.geometry]).plot(ax=ax, color=ZPAL[(r["생활권"] - 1) % 9], edgecolor="white", lw=1.5, zorder=1)
gpd.GeoSeries(pf.buffer(PROTECT_R).union_all().intersection(D)).plot(ax=ax, facecolor="none", edgecolor="#9a9a9a", hatch="////", lw=0, alpha=0.5, zorder=2)
gpd.GeoSeries([D]).boundary.plot(ax=ax, color=INK, lw=1.2, zorder=3)
for z in range(1, NC + 1):
    g = hh[hh["생활권"] == z]
    ax.scatter(g.geometry.x, g.geometry.y, s=np.sqrt(g["세대"]) * 5, color=ZDOT[(z - 1) % 9], lw=0, alpha=0.8, zorder=3)
for _, r in Z.iterrows():
    g = hh[hh["생활권"] == r["생활권"]]
    p = Point(np.average(g.geometry.x, weights=g["세대"]), np.average(g.geometry.y, weights=g["세대"]))
    if not r.geometry.contains(p):
        p = r.geometry.representative_point()
    # 거점·후보 표시와 겹치지 않으면서 중심에 가까운 자리 (생활권 안 15m 격자에서 고름)
    marks = list(H.geometry) + list(A.geometry) + list(C.geometry)
    inner = r.geometry.buffer(-35)
    x0, y0, x1, y1 = r.geometry.bounds
    grid = [Point(x, y) for x in np.arange(x0, x1, 15) for y in np.arange(y0, y1, 15)]
    grid = [q for q in grid if inner.contains(q)] or [p]
    p = max(grid, key=lambda q: min(min(q.distance(m) for m in marks), 110) - 0.35 * q.distance(p))
    ax.annotate(f"생활권 {r['생활권']}\n{r['세대']:,}세대", (p.x, p.y), ha="center", va="center", fontsize=8.5, color=INK, zorder=9,
                bbox=dict(boxstyle="round,pad=0.25", fc="white", ec="none", alpha=0.8))
for _, r in C.iterrows():
    hub = r["역할"] == "거점"
    ax.scatter(r.geometry.x, r.geometry.y, s=230 if hub else 120, color=GREEN if hub else "white", edgecolor=GREEN if hub else "#8a8a8a", lw=1.6, zorder=6)
    ax.annotate(str(int(r["번호"])), (r.geometry.x, r.geometry.y), ha="center", va="center", fontsize=8, fontweight="bold",
                color="white" if hub else "#6a6a6a", zorder=7)
for _, r in A[A["용도"] == "대안"].iterrows():
    first = r["순위"] == 1
    ax.scatter(r.geometry.x, r.geometry.y, s=380 if first else 110, marker="*", color=ORANGE if first else "white",
               edgecolor=ORANGE, lw=1.1, zorder=8)
    ax.annotate(r["id"], (r.geometry.x, r.geometry.y), xytext=(8, 5) if first else (-4, -13), textcoords="offset points", fontsize=8.5 if first else 7.5,
                color=ORANGE, fontweight="bold" if first else "normal", zorder=9)
bx = D.bounds
ax.set_xlim(bx[0] - 40, bx[2] + 40)
ax.set_ylim(bx[1] - 40, bx[3] + 40)
ax.legend(handles=[
    Line2D([], [], marker="o", ls="", ms=12, color=GREEN, label="거점: 생활권 안 1차 후보"),
    Line2D([], [], marker="*", ls="", ms=16, color=ORANGE, label="거점: 대안 지점 1순위 (실사 필요)"),
    Line2D([], [], marker="*", ls="", ms=10, mfc="white", mec=ORANGE, label="대안 2·3순위"),
    Line2D([], [], marker="o", ls="", ms=9, mfc="white", mec="#8a8a8a", label="예비 1차 후보 (같은 생활권 2번째 이후)"),
    Patch(facecolor="white", edgecolor="#9a9a9a", hatch="////", label=f"보호구역 대상시설 {PROTECT_R}m"),
    Line2D([], [], marker="o", ls="", ms=6, color=ZDOT[0], label="세대 필지 (크기 = 세대 수)"),
], loc="lower left", fontsize=8.5, frameon=True, framealpha=0.95)
ax.set_title(f"신월1동 도로 생활권 {NC}곳 = 그린존 거점 {NC}곳\n1차 후보 {int((~H['실사필요']).sum())}곳 + 대안 지점 {int(H['실사필요'].sum())}곳 (Louvain, 아파트 제외 {HS['총세대']:,}세대)",
             loc="left", fontsize=12, color=INK)
ax.set_axis_off()
ax.set_aspect("equal")
plt.tight_layout()
plt.savefig(os.path.join(OUT, "greenzone2nd_map.png"), dpi=160)
plt.close()

# ---- 민감도 그림 ---------------------------------------------------------------------------------
fig, (a1, a2) = plt.subplots(1, 2, figsize=(11, 4.2))
for g, col, ls in zip(SENS_G, ["#9bbfe9", BLUE, "#174a87"], ["--", "-", ":"]):
    s = SENS[SENS["감마"] == g]
    a1.plot(s["연결반경_m"], s["생활권수"], ls, marker="o", color=col, lw=2 if g == GAMMA else 1.4, ms=5, label=f"γ = {g:g}")
a1.axvline(LINK_R, color=ORANGE, lw=1, alpha=0.7)
a1.annotate(f"본 분석\n{LINK_R}m · γ {GAMMA:g} → {META['생활권수']}곳", (LINK_R, META["생활권수"]), xytext=(18, 80), textcoords="offset points",
            fontsize=9, color=ORANGE, arrowprops=dict(arrowstyle="-", color=ORANGE, lw=0.8))
a1.set_xlabel("필지 연결 반경 (도로거리, m)", color=MUTED)
a1.set_ylabel("생활권 수 (= 거점 수)", color=MUTED)
a1.set_title("반경 150m부터 생활권 수가 7~12곳에서 안정", loc="left", fontsize=10.5, color=INK)
a1.legend(frameon=False, fontsize=9)
s = SENS[SENS["감마"] == GAMMA]
a2.plot(s["연결반경_m"], s["이상중심_300m안"] * 100, "-o", color=BLUE, ms=5, lw=2, label="300m 안")
a2.plot(s["연결반경_m"], s["이상중심_200m안"] * 100, "--o", color="#9bbfe9", ms=5, lw=1.4, label="200m 안")
a2.axvline(LINK_R, color=ORANGE, lw=1, alpha=0.7)
a2.set_ylim(0, 105)
a2.set_xlabel("필지 연결 반경 (도로거리, m)", color=MUTED)
a2.set_ylabel("세대 비율 (%)", color=MUTED)
a2.set_title("생활권마다 이상적 중심에 거점을 둘 때 보행 접근성 (γ=1)", loc="left", fontsize=10.5, color=INK)
a2.legend(frameon=False, fontsize=9, loc="lower left")
for a in (a1, a2):
    a.grid(axis="y", color=GRID)
    a.spines[["top", "right"]].set_visible(False)
    a.set_xticks(SENS_R)
plt.tight_layout()
plt.savefig(os.path.join(OUT, "greenzone2nd_sensitivity.png"), dpi=170)
plt.close()

# ---- 엑셀 ----------------------------------------------------------------------------------------
FONT = "맑은 고딕"
HF = PatternFill("solid", fgColor="1B6B4A")
BOR = Border(*[Side(style="thin", color="CCCCCC")] * 4)
TINT = {"1차 후보": "E3F4EC", "대안 지점 (실사 필요)": "FDEBDD", "거점": "E3F4EC", "예비": "F2F2F2", "대안": "FDEBDD", "참고(후보 있음)": "EAF1FB"}
wb = Workbook()
wb.remove(wb.active)


def sheet(name, df, title, note, pct=(), m0=(), widths=None, tint=None, start=4):
    ws = wb.create_sheet(name)
    ws.cell(1, 1, title).font = Font(name=FONT, size=13, bold=True, color="1B6B4A")
    ws.cell(2, 1, note).font = Font(name=FONT, size=9, color="666666")
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(start, j, c)
        x.font, x.fill, x.border = Font(name=FONT, bold=True, color="FFFFFF", size=10), HF, BOR
        x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start].height = 32
    for i, row in enumerate(df.itertuples(index=False), start + 1):
        t = TINT.get(df.iloc[i - start - 1][tint]) if tint else None
        for j, v in enumerate(row, 1):
            if isinstance(v, (float, np.floating)) and np.isnan(v):
                v = None
            if isinstance(v, (np.integer,)):
                v = int(v)
            if isinstance(v, (np.floating,)):
                v = float(v)
            if isinstance(v, (np.bool_,)):
                v = bool(v)
            x = ws.cell(i, j, v)
            x.font, x.border = Font(name=FONT, size=10), BOR
            col = df.columns[j - 1]
            x.number_format = "0%" if col in pct else "#,##0" if col in m0 else "General"
            if t:
                x.fill = PatternFill("solid", fgColor=t)
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(c, max(9, min(30, len(str(c)) * 1.8 + 2)))
    ws.freeze_panes = ws.cell(start + 1, 2)
    return ws


# 요약
ws = wb.create_sheet("요약")
ws["A1"] = "그린존 2차: 도로 생활권(Louvain) 기반 거점 배치"
ws["A1"].font = Font(name=FONT, size=14, bold=True, color="1B6B4A")
near = CMP.iloc[1]
lines = [
    ("대상", f"신월1동 아파트 외 주거 {HS['필지수']:,}필지 · {HS['총세대']:,}세대 (건축물대장 표제부)"),
    ("생활권 수 = 거점 수", f"{NC}곳 (Louvain, 연결 반경 {LINK_R}m, γ {GAMMA:g}, modularity {META['modularity']:.3f})"),
    ("1차 후보로 정한 거점", f"{int((~H['실사필요']).sum())}곳: " + ", ".join(H.loc[~H["실사필요"], "이름"])),
    ("대안 지점 (실사)", f"{int(H['실사필요'].sum())}곳: 생활권 " + ", ".join(str(z) for z in H.loc[H["실사필요"], "생활권"]) + " (각 1~3순위, '대안 지점' 시트)"),
    ("예비 1차 후보", ", ".join(f"{int(r['번호'])}번" for _, r in cand[cand["역할"] == "예비"].iterrows()) + " (같은 생활권에 더 나은 후보가 있음)"),
    ("평균 보행거리", f"{near['평균보행_m']:.0f}m (가장 가까운 거점 이용 시)"),
    ("300m 안 세대", f"{near['300m안']:.0%} · 400m 안 {near['400m안']:.0%}"),
    ("수거 순회 거리", f"{S['route_len']/1000:.1f}km (거점 {NC}곳 한 바퀴, 도로 기준)"),
]
for i, (k, v) in enumerate(lines, 3):
    ws.cell(i, 1, k).font = Font(name=FONT, size=10, bold=True)
    ws.cell(i, 2, v).font = Font(name=FONT, size=10)
ws.column_dimensions["A"].width = 20
ws.column_dimensions["B"].width = 110
r0 = len(lines) + 5
ws.cell(r0, 1, "시트 안내").font = Font(name=FONT, size=11, bold=True, color="1B6B4A")
guide = [("거점", f"생활권별 거점 {NC}곳과 보행 접근성·담당 세대"), ("1차 후보 연결", "11곳이 어느 생활권에 속하고 거점/예비 중 무엇인지"),
         ("대안 지점", "후보가 없는 생활권의 실사 대상 1~3순위 + 후보가 있는 생활권의 최적 도로 지점(참고)"),
         ("안 비교", "생활권 거점 vs 1차 후보만"), ("Louvain 민감도", "설정값별 생활권 수"), ("세대 정제", "표제부 → 세대 필지 정제 과정"), ("가정", "바꿀 수 있는 기준값")]
for i, (k, v) in enumerate(guide, r0 + 1):
    ws.cell(i, 1, k).font = Font(name=FONT, size=10, bold=True)
    ws.cell(i, 2, v).font = Font(name=FONT, size=10)

Hx = H.drop(columns="geometry").copy()
ws = sheet("거점", Hx, f"생활권별 거점 {NC}곳", "보행거리는 '자기 생활권 거점'까지 도로 기준. 초록 = 1차 후보, 주황 = 대안 지점(실사 필요)",
           pct=("200m안", "300m안", "400m안"), m0=("담당세대", "담당필지", "평균보행_m", "세대90%_m", "최대보행_m", "보호시설까지_m", "생활권_최적도로지점_평균보행_m"),
           widths={"이름": 40, "구분": 18, "생활권안_1차후보": 14, "접한도로": 14}, tint="구분")
n = len(Hx) + 5
ws.cell(n, 1, "합계").font = Font(name=FONT, bold=True)
ws.cell(n, 6, f"=SUM(F5:F{n-1})").number_format = "#,##0"
ws.cell(n, 7, f"=SUM(G5:G{n-1})").number_format = "#,##0"
ws.cell(n, 8, f"=SUMPRODUCT(F5:F{n-1},H5:H{n-1})/F{n}").number_format = "#,##0"
for col in "KLM":
    ws[f"{col}{n}"] = f"=SUMPRODUCT($F5:$F{n-1},{col}5:{col}{n-1})/$F{n}"
    ws[f"{col}{n}"].number_format = "0%"
for c in range(1, 14):
    ws.cell(n, c).font = Font(name=FONT, bold=True, size=10)

sheet("1차 후보 연결", C.drop(columns="geometry"), "1차 후보 11곳 → 생활권 연결",
      "후보 위치가 생활권 영역 안에 있으면 연결. 한 생활권에 2곳 이상이면 그 생활권 세대의 평균 보행거리가 짧은 1곳 = 거점, 나머지 = 예비",
      pct=("생활권_300m안",), m0=("생활권_평균보행_m", "생활권_최적도로지점_평균보행_m", "보호시설과의 거리_m"), widths={"이름": 34}, tint="역할")
sheet("대안 지점", A.drop(columns="geometry"), "대안 지점: 후보가 없는 생활권의 실사 대상",
      f"조건: 등록 도로폭 {ROAD_MIN:g}m 이상 + 보호구역 대상시설 {PROTECT_R}m 밖인 도로 위 지점({STEP}m 간격). 그 생활권 세대 평균 보행거리 순, 1~3순위는 서로 {ALT_SPACING}m 이상 떨어짐. 파랑 = 후보 있는 생활권의 최적 지점(참고용)",
      pct=("생활권_300m안",), m0=("보호시설까지_m", "교차로까지_m", "생활권_평균보행_m"), widths={"가까운_건물주소": 30, "최근접_보호시설": 28}, tint="용도")
CM2 = CMP.copy()
CM2.loc[len(CM2)] = None
sheet("안 비교", CMP, "거점 배치안 비교 (아파트 제외 세대, 도로 보행거리)",
      f"수거 순회: 생활권 거점 {NC}곳 {S['route_len']/1000:.2f}km, 1차 후보 11곳 {S['cand_route']/1000:.2f}km",
      pct=("200m안", "300m안", "400m안"), m0=("평균보행_m", "세대90퍼_m", "최대보행_m"), widths={"안": 40})
sheet("Louvain 민감도", SENS, "설정값을 바꾸면 생활권 수가 어떻게 변하나",
      f"본 분석 = 반경 {LINK_R}m, γ {GAMMA:g}. 시드 {len(SEEDS)}번 중 modularity 최고 결과. 이상중심 = 생활권마다 세대 보행거리 합이 최소인 필지",
      pct=("이상중심_200m안", "이상중심_300m안"), m0=("이상중심_평균보행_m",))
steps = pd.DataFrame([
    ("표제부 전체", HS["표제부_행수"], "양천구 신월동(법정동) = 신월1~7동"),
    ("신월1동 건물에 연결", HS["신월1동_표제부_행수"], f"지번 → 전자지도 건물, 신월1동 경계 안 (주거 건물 연결률 {HS['연결률']:.1%})"),
    ("아파트 제외", -HS["아파트_행수_신월1동"], f"기타용도 '아파트' {HS['아파트_행수_신월1동']}동, 해당 필지 {HS['아파트_필지수_신월1동']}곳 전체 제외"),
    ("세대가 있는 건물", HS["세대있는_건물수"], "세대수 + 가구수 > 0"),
    ("  └ 상가주택 포함", HS["상가주택_건물수"], f"주용도 근린생활시설 등이지만 세대·가구 있음 ({HS['상가주택_세대']}세대)"),
    ("  └ 0 → 1세대 보정", HS["보정1세대_건물수"], "주용도가 주택인데 세대·가구 0"),
    ("세대 필지 (생활권 점)", HS["필지수"], f"필지 단위로 합침, 총 {HS['총세대']:,}세대 · 필지당 중앙값 {HS['필지당_세대_중앙값']:.0f}세대"),
], columns=["단계", "건수", "설명"])
sheet("세대 정제", steps, "건축물대장 표제부 → 세대 필지", "", m0=("건수",), widths={"단계": 22, "설명": 70})
ht = pd.DataFrame([(k, v, v / HS["총세대"]) for k, v in HS["주택유형별_세대"].items()], columns=["주택유형", "세대", "비율"])
ws = wb["세대 정제"]
ws.cell(14, 1, "주택유형별 세대").font = Font(name=FONT, size=11, bold=True, color="1B6B4A")
for j, c in enumerate(ht.columns, 1):
    x = ws.cell(15, j, c)
    x.font, x.fill, x.border = Font(name=FONT, bold=True, color="FFFFFF", size=10), HF, BOR
for i, r in enumerate(ht.itertuples(index=False), 16):
    for j, v in enumerate(r, 1):
        x = ws.cell(i, j, float(v) if isinstance(v, (np.floating, float)) else int(v) if isinstance(v, (np.integer, int)) else v)
        x.font, x.border = Font(name=FONT, size=10), BOR
        x.number_format = "0%" if j == 3 else "#,##0"
AS = pd.DataFrame([
    ("LINK_R", LINK_R, "m", "도로로 이 거리 안의 세대 필지끼리 연결 (생활권 크기를 정하는 값)"),
    ("DECAY", DECAY, "m", "연결 강도 = exp(−거리/DECAY) × √(세대i×세대j)"),
    ("GAMMA", GAMMA, "", "Louvain 해상도. 1 = 표준 modularity (조정하지 않음)"),
    ("SEEDS", len(SEEDS), "번", "Louvain 반복 횟수 (modularity 최고 결과 사용)"),
    ("ROAD_MIN", ROAD_MIN, "m", "대안 지점 등록 도로폭 하한 (5t 수거차 양방향)"),
    ("PROTECT_R", PROTECT_R, "m", "대안 지점 보호구역 대상시설 반경"),
    ("STEP", STEP, "m", "대안 지점을 찾을 때 도로를 훑는 간격"),
    ("ALT_SPACING", ALT_SPACING, "m", "대안 1~3순위끼리 최소 간격"),
], columns=["이름 (src/gzlib.py)", "값", "단위", "의미"])
sheet("가정", AS, "기준값", "src/gzlib.py 맨 위에서 바꾸고 python run_all.py 로 다시 실행", widths={"이름 (src/gzlib.py)": 20, "의미": 60})
wb.save(os.path.join(OUT, "greenzone2nd.xlsx"))

# ---- html 지도 ------------------------------------------------------------------------------------
def gj(g):
    return json.loads(g.to_crs(4326).to_json(drop_id=True))


prot = gpd.GeoDataFrame(pf[["시설명", "구분"]], geometry=pf.buffer(PROTECT_R), crs=CRS_M)
data = dict(zones=gj(Z), hubs=gj(H.drop(columns=["실사필요"]).assign(실사=H["실사필요"].map({True: "필요", False: "-"}))),
            cands=gj(C), alts=gj(A), hh=gj(hh[["세대", "주택유형", "생활권", "geometry"]]), prot=gj(prot),
            dong=gj(gpd.GeoDataFrame(geometry=[D], crs=CRS_M)))
html = """<!doctype html><html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>그린존 2차 거점 지도</title>
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<style>html,body,#map{height:100%;margin:0;font-family:'Malgun Gothic','Apple SD Gothic Neo',sans-serif}
.box{background:#fff;padding:8px 10px;border-radius:6px;box-shadow:0 1px 4px rgba(0,0,0,.25);font-size:12px;line-height:1.6;max-width:260px}
.box b{color:#1B6B4A}.lab{background:rgba(255,255,255,.85);border:none;box-shadow:none;font-weight:600;font-size:11px;color:#222}
.star{color:#eb6834;font-size:22px;line-height:22px;text-shadow:0 0 2px #fff}.star.s{font-size:15px}
.num{background:#1baf7a;color:#fff;border-radius:50%;width:22px;height:22px;text-align:center;line-height:22px;font-weight:700;font-size:11px;border:2px solid #fff}
.num.r{background:#fff;color:#777;border:2px solid #999;width:18px;height:18px;line-height:18px}
table.p td{padding:1px 6px 1px 0;vertical-align:top}</style></head><body><div id="map"></div><script>
const D=__DATA__;const PAL=__PAL__;
const map=L.map('map');L.tileLayer('https://{s}.tile.openstreetmap.org/{z}/{x}/{y}.png',{maxZoom:19,attribution:'&copy; OpenStreetMap'}).addTo(map);
const tbl=o=>'<table class="p">'+Object.entries(o).filter(([k,v])=>v!==null&&v!==undefined&&v!=='').map(([k,v])=>`<tr><td>${k}</td><td><b>${typeof v==='number'?(/^(위도|경도)$/.test(k)?v:Math.abs(v)<=1&&/안$/.test(k)?Math.round(v*100)+'%':Math.round(v*10)/10):v}</b></td></tr>`).join('')+'</table>';
const zones=L.geoJSON(D.zones,{style:f=>({color:'#fff',weight:2,fillColor:PAL[(f.properties['생활권']-1)%9],fillOpacity:.45}),
 onEachFeature:(f,l)=>{l.bindPopup('<b>생활권 '+f.properties['생활권']+'</b>'+tbl(f.properties));l.bindTooltip('생활권 '+f.properties['생활권'],{permanent:true,direction:'center',className:'lab'})}}).addTo(map);
const prot=L.geoJSON(D.prot,{style:{color:'#888',weight:1,dashArray:'4 4',fillColor:'#888',fillOpacity:.08},onEachFeature:(f,l)=>l.bindPopup(f.properties['시설명']+' 200m')});
const hh=L.geoJSON(D.hh,{pointToLayer:(f,ll)=>L.circleMarker(ll,{radius:Math.sqrt(f.properties['세대'])*1.3+1,color:'#555',weight:.3,fillColor:PAL[(f.properties['생활권']-1)%9],fillOpacity:.9}),
 onEachFeature:(f,l)=>l.bindPopup(tbl(f.properties))}).addTo(map);
const cands=L.geoJSON(D.cands,{filter:f=>f.properties['역할']!=='거점',pointToLayer:(f,ll)=>L.marker(ll,{icon:L.divIcon({className:'',html:'<div class="num r">'+f.properties['번호']+'</div>',iconSize:[22,22]})}),
 onEachFeature:(f,l)=>l.bindPopup('<b>예비 후보 '+f.properties['번호']+'번</b> '+f.properties['이름']+tbl({생활권:f.properties['생활권'],생활권_평균보행_m:f.properties['생활권_평균보행_m'],생활권_300m안:f.properties['생활권_300m안']}))}).addTo(map);
const alts=L.geoJSON(D.alts,{filter:f=>f.properties['용도']==='대안'&&f.properties['순위']>1,pointToLayer:(f,ll)=>L.marker(ll,{icon:L.divIcon({className:'',html:'<div class="star s">☆</div>',iconSize:[16,16]})}),
 onEachFeature:(f,l)=>l.bindPopup('<b>대안 '+f.properties.id+'</b>'+tbl(f.properties))}).addTo(map);
const hubs=L.geoJSON(D.hubs,{pointToLayer:(f,ll)=>{const a=f.properties['거점id'].startsWith('A');return L.marker(ll,{zIndexOffset:1000,icon:L.divIcon({className:'',html:a?'<div class="star">★</div>':'<div class="num">'+f.properties['이름'].split('번')[0]+'</div>',iconSize:[24,24]})})},
 onEachFeature:(f,l)=>l.bindPopup('<b>생활권 '+f.properties['생활권']+' 거점</b><br>'+f.properties['이름']+tbl(f.properties))}).addTo(map);
L.geoJSON(D.dong,{style:{color:'#222',weight:2,fill:false}}).addTo(map);
map.fitBounds(zones.getBounds(),{padding:[10,10]});
L.control.layers(null,{'생활권':zones,'세대 필지':hh,'거점 9곳':hubs,'예비 1차 후보':cands,'대안 2·3순위':alts,'보호구역 200m':prot},{collapsed:false}).addTo(map);
const lg=L.control({position:'bottomleft'});lg.onAdd=()=>{const d=L.DomUtil.create('div','box');d.innerHTML='<b>그린존 2차: 생활권 __N__곳 = 거점 __N__곳</b><br><span style="color:#1baf7a">●</span> 거점(1차 후보) <span style="color:#eb6834">★</span> 거점(대안 지점·실사)<br><span style="color:#999">○</span> 예비 후보 <span style="color:#eb6834">☆</span> 대안 2·3순위<br>점 크기 = 세대 수, 색 = 생활권';return d};lg.addTo(map);
</script></body></html>"""
html = html.replace("__DATA__", json.dumps(data, ensure_ascii=False)).replace("__PAL__", json.dumps(ZDOT)).replace("__N__", str(NC))
open(os.path.join(OUT, "greenzone2nd_map.html"), "w", encoding="utf-8").write(html)
print("결과물:", sorted(os.listdir(OUT)))
